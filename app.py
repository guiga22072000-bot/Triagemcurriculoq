import os
import re
import json
import io
import uuid
from threading import Thread, Lock
from concurrent.futures import ThreadPoolExecutor, as_completed

from flask import Flask, render_template, request, send_file, redirect, flash, jsonify
from werkzeug.utils import secure_filename

import openpyxl
import pdfplumber
import docx
from openai import OpenAI
from supabase import create_client

# ========================
# CONFIG
# ========================
app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-key")

UPLOAD_FOLDER = "uploads"
ALLOWED_EXTENSIONS = {"pdf", "docx", "doc", "txt"}
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY", "")
client = OpenAI(api_key=OPENAI_API_KEY) if OPENAI_API_KEY else None

# ✅ SUPABASE (configure depois)
SUPABASE_URL = "https://djitgqkgypkjfhluqgrd.supabase.co"
SUPABASE_KEY = "sb_publishable_eJgqpcF1yCDdvweUde5LZg_L2CbXqT_"
supabase = create_client(SUPABASE_URL, SUPABASE_KEY) if SUPABASE_URL else None

MAX_WORKERS = 2

tasks = {}
tasks_lock = Lock()

# ========================
# NORMALIZAÇÃO TEXTO 🔥
# ========================
def normalize_text(text):
    text = re.sub(r"(linkedin\.com/in/)\s*\n\s*", r"\1", text)
    text = re.sub(r"(https?://[^\s]+)\s*\n\s*([^\s]+)", r"\1\2", text)
    return text

# ========================
# EXTRAÇÃO
# ========================
def extract_text_from_pdf(filepath):
    text = ""
    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages[:10]:
                t = page.extract_text()
                if t:
                    text += t + "\n"
    except:
        pass
    return text


def extract_text_from_docx(filepath):
    text = ""
    try:
        doc = docx.Document(filepath)
        for para in doc.paragraphs[:300]:
            text += para.text + "\n"
    except:
        pass
    return text


def extract_text(filepath, ext):
    if ext == "pdf":
        return extract_text_from_pdf(filepath)
    elif ext in ("doc", "docx"):
        return extract_text_from_docx(filepath)
    return ""


# ========================
# FALLBACK LINKEDIN 🔥
# ========================
def extract_linkedin(text):
    text = normalize_text(text)

    match = re.search(r"(https?://)?(www\.)?linkedin\.com/in/[A-Za-z0-9\-_%]+", text)

    if not match:
        return "Não encontrado"

    url = match.group(0)

    if not url.startswith("http"):
        url = "https://" + url

    if not re.search(r"linkedin\.com/in/[A-Za-z0-9\-_%]{3,}", url):
        return "Não encontrado"

    return url


def regex_fallback_extract(text):
    email = re.search(r"[^@\s]+@[^@\s]+", text)
    phone = re.search(r"\(?\d{2}\)?\s?\d{4,5}-?\d{4}", text)

    return {
        "nome": text.split("\n")[0][:80],
        "whatsapp": phone.group(0) if phone else "Não encontrado",
        "email": email.group(0) if email else "Não encontrado",
        "linkedin": extract_linkedin(text)
    }

# ========================
# IA
# ========================
def analyze_resume(text, vaga):
    text = normalize_text(text)

    if not client:
        fallback = regex_fallback_extract(text)
        fallback.update({"score": "-", "status": "-", "justificativa": "Sem API"})
        return fallback

    prompt = f"""
Você é analista de RH.

VAGA:
{vaga}

CURRÍCULO:
{text[:12000]}

Retorne JSON:

{{
"nome":"",
"whatsapp":"",
"email":"",
"linkedin":"",
"score":"",
"status":"",
"justificativa":""
}}

Regras:
- Se LinkedIn estiver em duas linhas, reconstruir
- Se estiver como /in/ apenas → ignorar
"""

    try:
        res = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2
        )

        data = json.loads(res.choices[0].message.content)

        # fallback forte de linkedin
        if not data.get("linkedin") or "linkedin.com/in/" not in data.get("linkedin"):
            data["linkedin"] = extract_linkedin(text)

        return data

    except:
        fallback = regex_fallback_extract(text)
        fallback.update({"score": "-", "status": "Erro IA"})
        return fallback

# ========================
# PROCESSAMENTO
# ========================
def background_task(task_id, files, vaga, job_title):

    results = []

    def process(path):
        ext = path.split(".")[-1]
        text = extract_text(path, ext)
        os.remove(path)

        result = analyze_resume(text, vaga)
        result["arquivo"] = os.path.basename(path)
        result["vaga"] = job_title

        return result

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futures = [ex.submit(process, f) for f in files]

        for i, future in enumerate(as_completed(futures)):
            results.append(future.result())

            with tasks_lock:
                tasks[task_id]["progress"] = int((i+1)/len(files)*100)

    with tasks_lock:
        tasks[task_id]["done"] = True
        tasks[task_id]["results"] = results


# ========================
# ROTAS
# ========================
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/processar", methods=["POST"])
def processar():

    vaga = request.form.get("job_profile")
    job_title = request.form.get("job_title")

    files = request.files.getlist("resumes")

    task_id = str(uuid.uuid4())
    paths = []

    for f in files:
        filename = secure_filename(f.filename)
        path = os.path.join(UPLOAD_FOLDER, f"{task_id}_{filename}")
        f.save(path)
        paths.append(path)

    with tasks_lock:
        tasks[task_id] = {"progress": 0, "done": False, "results": []}

    Thread(target=background_task, args=(task_id, paths, vaga, job_title)).start()

    return redirect(f"/status/{task_id}")


@app.route("/status/<task_id>")
def status(task_id):
    return render_template("status.html", task_id=task_id)


@app.route("/progresso/<task_id>")
def progresso(task_id):
    return jsonify(tasks.get(task_id, {}))


@app.route("/resultado/<task_id>")
def resultado(task_id):
    task = tasks.get(task_id)

    if not task or not task["done"]:
        return redirect(f"/status/{task_id}")

    return render_template("resultado.html", results=task["results"])


# ========================
# SALVAR NO BANCO
# ========================
@app.route("/salvar_banco", methods=["POST"])
def salvar_banco():

    if not supabase:
        return "Configure o Supabase"

    results = json.loads(request.form.get("results_json"))

    for r in results:
        supabase.table("candidatos").insert({
            "nome": r.get("nome"),
            "whatsapp": r.get("whatsapp"),
            "email": r.get("email"),
            "linkedin": r.get("linkedin"),
            "vaga": r.get("vaga"),
            "score": r.get("score"),
            "status": r.get("status"),
            "justificativa": r.get("justificativa"),
            "arquivo": r.get("arquivo")
        }).execute()

    return redirect("/banco")


# ========================
# BANCO
# ========================
@app.route("/banco")
def banco():

    vaga = request.args.get("vaga")

    query = supabase.table("candidatos").select("*")

    if vaga:
        query = query.eq("vaga", vaga)

    data = query.execute()

    return render_template("banco.html", candidatos=data.data)


# ========================
# EXPORTAÇÃO
# ========================
def gerar_excel(results):

    wb = openpyxl.Workbook()
    ws = wb.active

    headers = ["Vaga", "Nome", "WhatsApp", "Email", "LinkedIn", "Score", "Status", "Justificativa"]
    ws.append(headers)

    for r in results:
        ws.append([
            r.get("vaga"),
            r.get("nome"),
            r.get("whatsapp"),
            r.get("email"),
            r.get("linkedin"),
            r.get("score"),
            r.get("status"),
            r.get("justificativa")
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return output


@app.route("/exportar", methods=["POST"])
def exportar():
    results = json.loads(request.form.get("results_json"))
    output = gerar_excel(results)

    return send_file(output, as_attachment=True, download_name="resultado.xlsx")


@app.route("/exportar_banco", methods=["POST"])
def exportar_banco():

    vaga = request.form.get("vaga")

    query = supabase.table("candidatos").select("*")

    if vaga:
        query = query.eq("vaga", vaga)

    data = query.execute()

    output = gerar_excel(data.data)

    return send_file(output, as_attachment=True, download_name="banco.xlsx")


# ========================
if __name__ == "__main__":
    app.run(debug=True)
